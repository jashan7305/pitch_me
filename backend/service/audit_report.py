from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt

from models.audit import AuditReport


def _add_heading(
    document: Document,
    text: str,
    level: int = 1,
) -> None:
    document.add_heading(text, level=level)


def _add_label_value(
    document: Document,
    label: str,
    value: str,
) -> None:
    paragraph = document.add_paragraph()

    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True

    paragraph.add_run(value)


def _add_status(
    document: Document,
    status: str,
) -> None:
    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        status.replace("_", " ").upper()
    )
    run.bold = True
    run.font.size = Pt(12)


def _add_evidence(
    document: Document,
    evidence: str,
) -> None:
    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        f'"{evidence}"'
    )
    run.italic = True


def _add_claim_section(
    document: Document,
    claim_number: int,
    claim,
) -> None:
    _add_heading(
        document,
        f"Claim {claim_number}",
        level=2,
    )

    _add_label_value(
        document,
        "Claim",
        claim.claim,
    )

    _add_label_value(
        document,
        "Claim Type",
        claim.claim_type,
    )

    _add_label_value(
        document,
        "Status",
        claim.status,
    )

    _add_label_value(
        document,
        "Confidence",
        f"{claim.confidence:.0%}",
    )

    if claim.explanation:
        _add_label_value(
            document,
            "Auditor Explanation",
            claim.explanation,
        )

    if claim.source_document:
        _add_label_value(
            document,
            "Source Document",
            claim.source_document,
        )

    if claim.source_page:
        _add_label_value(
            document,
            "Source Page",
            str(claim.source_page),
        )

    if claim.source_chunk_id:
        _add_label_value(
            document,
            "Source Chunk",
            claim.source_chunk_id,
        )

    if claim.evidence:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("Evidence")
        run.bold = True

        _add_evidence(
            document,
            claim.evidence,
        )

    if claim.status != "supported":
        paragraph = document.add_paragraph()

        run = paragraph.add_run(
            "Action: Review this claim before presenting."
        )
        run.bold = True


def generate_audit_report(
    audit: AuditReport,
    output_path: str | Path,
) -> Path:
    """
    Generate a human-readable DOCX audit report.

    This function only formats the AuditReport. It does not
    perform or modify the audit.
    """

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document = Document()

    # ---------------------------------------------------------
    # Page setup
    # ---------------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # ---------------------------------------------------------
    # Title
    # ---------------------------------------------------------

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        "AI CLAIM AUDIT REPORT"
    )

    run.bold = True
    run.font.size = Pt(24)

    subtitle = document.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(
        "Client-Specific Insurance Pitch"
    )

    run.font.size = Pt(14)

    document.add_paragraph()

    # ---------------------------------------------------------
    # Client / policy information
    # ---------------------------------------------------------

    _add_heading(
        document,
        "Report Information",
        level=1,
    )

    _add_label_value(
        document,
        "Client",
        audit.company_name,
    )

    _add_label_value(
        document,
        "Policy Provider",
        audit.policy_provider,
    )

    if audit.policy_product:
        _add_label_value(
            document,
            "Policy Product",
            audit.policy_product,
        )

    # ---------------------------------------------------------
    # Summary
    # ---------------------------------------------------------

    _add_heading(
        document,
        "Audit Summary",
        level=1,
    )

    table = document.add_table(
        rows=1,
        cols=2,
    )

    table.style = "Table Grid"

    header = table.rows[0].cells

    header[0].text = "Metric"
    header[1].text = "Count"

    summary_rows = [
        ("Total Claims", audit.total_claims),
        ("Supported", audit.supported_claims),
        (
            "Partially Supported",
            audit.partially_supported_claims,
        ),
        ("Unsupported", audit.unsupported_claims),
    ]

    for label, value in summary_rows:
        cells = table.add_row().cells

        cells[0].text = label
        cells[1].text = str(value)

        for cell in cells:
            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

    document.add_paragraph()

    overall_status = (
        "PASS — No unsupported claims identified."
        if audit.audit_passed
        else
        "REVIEW REQUIRED — Unsupported claims identified."
    )

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        f"Overall Audit Status: {overall_status}"
    )

    run.bold = True

    # ---------------------------------------------------------
    # Claim-by-claim audit
    # ---------------------------------------------------------

    _add_heading(
        document,
        "Claim-by-Claim Audit",
        level=1,
    )

    if not audit.claims:
        document.add_paragraph(
            "No claims were generated in the pitch."
        )

    for index, claim in enumerate(
        audit.claims,
        start=1,
    ):
        _add_claim_section(
            document,
            index,
            claim,
        )

        if index != len(audit.claims):
            document.add_paragraph(
                "─" * 60
            )

    # ---------------------------------------------------------
    # Footer
    # ---------------------------------------------------------

    for section in document.sections:
        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.text = (
            "Generated from the AI pitch audit results."
        )

        paragraph.runs[0].font.size = Pt(9)

    document.save(
        str(output_path)
    )

    return output_path